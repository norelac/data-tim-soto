"""Buat lampiran Excel: sumber data (bertanda tautan), indikator, ruas jalan, simulasi anggaran, asumsi biaya."""
import csv, json, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(r"D:\tim soto")
D = ROOT / "data"
ds = json.loads((D / "semarangcheck_dataset.json").read_text(encoding="utf-8"))
sumber = json.loads((D / "sumber_ruas_jalan.json").read_text(encoding="utf-8"))
ind = list(csv.DictReader((D / "indikator_kecamatan.csv").open(encoding="utf-8-sig")))
ruas = list(csv.DictReader((D / "ruas_jalan_2022_lengkap.csv").open(encoding="utf-8-sig")))

HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2E5C8A")
LINK = Font(color="1F4E79", underline="single", size=10)

wb = Workbook()

def sheet(name, headers, rows, widths, link_cols=()):
    ws = wb.create_sheet(name) if wb.sheetnames != ["Sheet"] else wb.active
    ws.title = name
    ws.append(headers)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.font = HDR
        c.fill = FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    for r in rows:
        ws.append(r)
    for i in link_cols:
        for row in range(2, ws.max_row + 1):
            c = ws.cell(row=row, column=i)
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value
                c.font = LINK
    ws.freeze_panes = "A2"
    return ws

# 1. sumber data
src = [[
    "Kondisi ruas jalan per kecamatan (panjang, lebar, kondisi)", "Portal Data Jawa Tengah (CKAN)",
    "Kecamatan & ruas", 2022, "Google Sheets / XLSX",
    "https://data.jatengprov.go.id/dataset/?q=ruas+jalan+kecamatan+kota+semarang",
    "16 dataset, lihat lembar 'Sumber_Ruas_Jalan'", "Sudah diunduh",
], [
    "Jumlah penduduk, luas wilayah, kepadatan per kecamatan", "BPS Kota Semarang", "Kecamatan", 2025, "HTML (JSON tertanam)",
    "https://semarangkota.bps.go.id/id/statistics-table/2/NDgjMg%3D%3D/luas-wilayah--jumlah-penduduk--dan-kepadatan-penduduk.html",
    "16 kecamatan; total cocok 1.722.421 jiwa", "Sudah diunduh",
], [
    "Harga satuan pekerjaan jalan (aspal, agregat, galian)", "HSP Bina Marga Edisi I 2026 Kota Semarang", "Item pekerjaan", 2026, "CSV (111 item)",
    "https://data.jatengprov.go.id/dataset/harga-satuan-pekerjaan-hsp-konstruksi-bidang-bina-marga-edisi-ke-i-tahun-2026-kota-semarang",
    "kode dipakai: C.1, E.1, F.1, F.24, F.27", "Sudah diunduh",
], [
    "Garis batas kecamatan untuk peta", "GADM 4.1 level 3 (Indonesia)", "Kecamatan", "rilis 2022", "GeoJSON",
    "https://geodata.ucdavis.edu/gadm/gadm4.1/json/gadm41_IDN_3.json.zip",
    "difilter Kota Semarang -> 16 poligon", "Sudah diunduh",
], [
    "Titik lampu jalan (PJU) per kecamatan", "Portal Data Jawa Tengah", "Kecamatan", 2022, "XLSX (Google Drive)",
    "https://data.jatengprov.go.id/dataset/data-titik-lampu-per-kecamatan-kota-semarang-desember-2022",
    "butuh izin akses", "Gagal (401) - di luar lingkup",
], [
    "Kondisi jalan setelah 2022 / data tren", "-", "-", "-", "-", "-", "tidak tersedia publik", "Tidak ada",
], [
    "Koordinat ruas jalan (lat/long)", "-", "-", "-", "-", "-", "dataset hanya memuat kode & nama ruas", "Tidak ada",
], [
    "Realisasi anggaran jalan per kecamatan", "-", "-", "-", "-", "-",
    "yang tersedia hanya belanja operasional kecamatan, bukan belanja jalan", "Tidak ada",
]]
sheet("Sumber_Data",
      ["Kebutuhan data", "Sumber", "Level", "Tahun", "Format", "Tautan (klik)", "Catatan", "Status"],
      src, [42, 34, 14, 8, 20, 60, 40, 24], link_cols=(6,))

# 2. tautan 16 dataset ruas jalan
sheet("Sumber_Ruas_Jalan",
      ["Kecamatan", "Judul dataset", "Tautan berkas asli (klik)", "Tipe file", "Halaman dataset", "Jumlah ruas"],
      [[s["kecamatan"], s["judul"], s["resource_url"],
        "Google Sheets" if s["tipe_file"] == "sheets" else "Google Drive (xlsx)",
        f"https://data.jatengprov.go.id/dataset/{s['paket']}", s["baris"]] for s in sumber],
      [16, 52, 60, 22, 60, 11], link_cols=(3, 5))

# 3. indikator + ranking
cols = ["peringkat", "kecamatan", "skor", "i0_volume_kerusakan", "i1_persen_kerusakan",
        "i2_kerusakan_per_1000_penduduk", "i3_kerusakan_per_km2", "i4_sebaran_ruas_rusak",
        "ruas", "km_total", "km_baik", "km_sedang", "km_rusak_ringan", "km_rusak_berat",
        "km_kondisi_kosong", "km_ekuivalen_rusak", "km_perlu_penanganan", "penduduk", "luas_km2",
        "estimasi_biaya_penuh"]
label = {"peringkat": "Peringkat (skenario seimbang)", "kecamatan": "Kecamatan", "skor": "Skor SAW",
         "i0_volume_kerusakan": "Volume kerusakan (km ekuivalen)", "i1_persen_kerusakan": "Intensitas kerusakan (rasio)",
         "i2_kerusakan_per_1000_penduduk": "Kerusakan per 1.000 penduduk (km)",
         "i3_kerusakan_per_km2": "Kerusakan per km2 (km)", "i4_sebaran_ruas_rusak": "Sebaran ruas rusak (rasio)",
         "ruas": "Jumlah ruas", "km_total": "Total panjang (km)", "km_baik": "Baik (km)", "km_sedang": "Sedang (km)",
         "km_rusak_ringan": "Rusak ringan (km)", "km_rusak_berat": "Rusak berat (km)",
         "km_kondisi_kosong": "Tanpa keterangan (km)", "km_ekuivalen_rusak": "Km ekuivalen rusak",
         "km_perlu_penanganan": "Km perlu penanganan", "penduduk": "Penduduk (jiwa)",
         "luas_km2": "Luas wilayah (km2)", "estimasi_biaya_penuh": "Estimasi biaya penuh (Rp)"}
rows = []
for r in ind:
    rows.append([float(r[c]) if c not in ("kecamatan", "peringkat", "ruas") else
                 (int(r[c]) if c != "kecamatan" else r[c]) for c in cols])
sheet("Indikator_Ranking", [label[c] for c in cols], rows,
      [11, 16, 10, 14, 14, 16, 14, 14, 9, 11, 10, 10, 13, 12, 13, 13, 13, 12, 12, 16])

# 4. data ruas lengkap
rc = ["kecamatan", "kode_ruas", "nama_ruas", "panjang_m", "lebar_m", "luas_m2", "kondisi_raw", "kondisi",
      "bobot", "km_ekuivalen_rusak", "biaya_perbaikan", "sumber_url"]
sheet("Ruas_Jalan_2022", rc,
      [[r["kecamatan"], r["kode_ruas"], r["nama_ruas"], float(r["panjang_m"]), float(r["lebar_m"]),
        float(r["luas_m2"]), r["kondisi_raw"], r["kondisi"], float(r["bobot"]),
        round(float(r["panjang_m"]) / 1000 * float(r["bobot"]), 4), float(r["biaya_perbaikan"]),
        r["sumber_url"]] for r in ruas],
      [16, 14, 34, 11, 9, 11, 14, 14, 8, 15, 16, 60], link_cols=(12,))

# 5. simulasi anggaran per batas anggaran
sim_rows = []
for s in ds["simulasi"]:
    sim_rows.append([s["anggaran"], s["ruas"], s["km"], s["km_ekuivalen"], s["cakupan_persen"],
                     s["terpakai"], s["sisa"], "; ".join(f"{k} ({v} km)" for k, v in list(s["kecamatan"].items())[:8])])
sheet("Simulasi_Anggaran",
      ["Batas anggaran (Rp)", "Ruas ditangani", "Km ditangani", "Km ekuivalen tertangani",
       "Cakupan kerusakan (%)", "Anggaran terpakai (Rp)", "Sisa (Rp)", "Kecamatan terbanyak"],
      sim_rows, [18, 14, 12, 20, 18, 20, 16, 60])

# 6. asumsi & metode
a = ds["asumsi_biaya"]
met = [
    ["Harga satuan", "HSP Bina Marga Edisi I 2026 Kota Semarang", a["sumber_harga"], ""],
    ["Biaya overlay (rusak ringan) per m2", "Rp", a["overlay_per_m2"], "resap pengikat 0,5 l/m2 + AC-WC 4 cm"],
    ["Biaya rekonstruksi (rusak berat) per m2", "Rp", a["rekonstruksi_per_m2"],
     "galian 20 cm + agregat kelas A 15 cm + resap pengikat + AC-WC 4 cm + AC-BC 5 cm"],
    ["Densitas campuran beraspal", "ton/m3", a["densitas_ton_m3"], "asumsi teknis umum"],
    ["Bobot keparahan", "", "", "rusak berat 1,0 | rusak ringan 0,6 | sedang 0,3 | baik & tanpa keterangan 0"],
    ["Metode perankingan", "", "", "SAW: Vi = Σ wj · rij, normalisasi min-max antar kecamatan"],
    ["Metode simulasi anggaran", "", "", "greedy: urutkan ruas rusak menurut km ekuivalen per rupiah, pilih sampai anggaran habis"],
    ["Cakupan data", "", "", "hanya ruas berkondisi rusak ringan & rusak berat yang dihitung biayanya; kondisi sedang di luar hitungan"],
    ["Catatan tahun", "", "", "kondisi jalan 2022, penduduk 2025, harga satuan 2026 - angka biaya bersifat estimasi indikatif"],
]
sheet("Metode_Asumsi", ["Komponen", "Satuan", "Nilai", "Keterangan"], met, [34, 12, 24, 70])

wb.save(ROOT / "SemarangCheck_Data.xlsx")
print("saved:", ROOT / "SemarangCheck_Data.xlsx", "| sheets:", wb.sheetnames)
